# Builds the SME wording review document (PS-Work-Zone-Confirmation-wording.docx)
# from the three configs in src/ and the engine's built-in messages: every
# EN/FR string with blank "New English" / "New French" columns, the
# untranslated values (emails, links, page addresses), and the built-in
# messages as an optional appendix. Each row carries its location in the
# config, so the filled-in copy can be applied to the JSON (by hand or with
# Copilot). Never deployed.
#
# Rebuild it after the wording in src/ changes:
#   pip install python-docx
#   python deploy/make-wording-doc.py
# (optional arguments: <repo root> <output .docx>; the defaults are this repo
# and docs/PS-Work-Zone-Confirmation-wording.docx)
# With PSZ_SELFTEST=1 set, it only checks which keys it marks "Formatting
# allowed" (takes_formatting) and exits.
import json, re, sys, os, datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(HERE)
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(REPO, 'docs', 'PS-Work-Zone-Confirmation-wording.docx')
SRC = os.path.join(REPO, 'src')
load = lambda n: json.load(open(os.path.join(SRC, n), encoding='utf-8'))
form, access, res = load('ps-zone-attestation.json'), load('ps-zone-attestation.access.json'), load('ps-zone-attestation.resources.json')
js = open(os.path.join(SRC, 'ps-zone-attestation.js'), encoding='utf-8').read()
VERSION = re.search(r"var VERSION = '([^']+)'", js).group(1)

def is_pair(v):
    return isinstance(v, dict) and 'en' in v and set(v) <= {'en', 'fr'}

# ---------------------------------------------------------------- walkers
def walk(node, path, out):
    """every { en, fr } pair under node, with a readable path"""
    if is_pair(node):
        out.append((path, node.get('en', ''), node.get('fr', '')))
        return
    if isinstance(node, dict):
        for k, v in node.items():
            if k.startswith('$'):
                continue
            walk(v, path + [k], out)
    elif isinstance(node, list):
        for i, v in enumerate(node):
            tag = None
            if isinstance(v, dict):
                if 'id' in v: tag = 'field "%s"' % v['id'] if 'type' in v else 'section "%s"' % v['id']
                elif 'value' in v: tag = 'zone "%s"' % v['value']
                elif 'type' in v: tag = 'item %d (%s)' % (i + 1, v['type'])
                elif 'heading' in v: tag = 'group %d' % (i + 1)
            if not tag:
                tag = ('page %d' % (i + 1)) if path and path[-1] == 'pages' else str(i + 1)
            walk(v, path + [tag], out)

def clean(path):
    # drop the structural noise Copilot doesn't need to see twice
    p = [s for s in path if s not in ('pages', 'sections', 'fields', 'choices', 'items')]
    return ' › '.join(p)

STRINGS_WHERE = {
    'submit': 'Submit button',
    'savingTitle': 'Full-page "saving" overlay: its title',
    'assignLoading': 'While the areas load',
    'assignLoadError': 'When the areas can\'t be loaded',
    'assignRowRequired': 'Error on an area with no zone chosen (area had no current zone)',
    'assignRowConfirmRequired': 'Error on an area neither confirmed nor changed',
    'assignAllRequired': 'Error summary when some areas aren\'t answered',
    'assignCurrent': 'Small heading above an area\'s current zone',
    'assignConfirm': 'Checkbox to confirm an area\'s current zone',
    'assignChange': 'Button to change an area\'s zone',
    'assignKeep': 'Button to undo a change (keep the current zone)',
    'assignNewPlaceholder': 'Dropdown prompt after "Change zone"',
    'assignConfirmed': 'Status tag on an area: confirmed',
    'assignChanged': 'Status tag on an area: changed',
    'assignSelected': 'Status tag on an area: zone chosen (had none)',
    'assignProgress': 'Progress line ("2 of 3 areas complete")',
    'assignSomeDone': 'Note when some areas were already submitted earlier',
    'rowsPartial': 'Error after a save that only partly went through',
    'rowsFailed': 'Error after a save that failed',
}
def where_form(path):
    p = path
    j = '.'.join(p)
    fixed = {
        'form.title': 'Form title (top of the header)',
        'form.intro': 'Intro text under the title',
        'form.headerNote.text': 'Blue callout in the header ("changes will be reviewed")',
        'form.headerUser.label': 'Label above the user card in the header',
        'errorContact.subject': 'Subject line of the "email the form administrator" problem report',
        'submitConfirm.title': 'Confirmation dialog: title',
        'submitConfirm.message': 'Confirmation dialog: message',
        'submitConfirm.confirm': 'Confirmation dialog: confirm button',
        'submitConfirm.cancel': 'Confirmation dialog: go-back button',
        'confirmation.title': 'Thank-you screen: title',
        'confirmation.message': 'Thank-you screen: message',
    }
    if j in fixed: return fixed[j]
    if p[0] == 'strings': return STRINGS_WHERE.get(p[1], 'Message: ' + p[1])
    if 'field "accessNotice"' in p:
        last = p[-1] if p[-1] != 'text' or 'button' not in p else 'button'
        return {'title': 'Access-list callout (top of the form): title', 'text': 'Access-list callout: text',
                'button': 'Access-list callout: button label'}.get(last, 'Access-list callout: ' + p[-1])
    if 'field "changeNotice"' in p: return 'Bottom message, shown only once a zone has been changed'
    if 'field "contactNotice"' in p: return 'Bottom message: who to contact (keep the [email](mailto:…) link form)'
    if 'field "zones"' in p:
        z = [s for s in p if s.startswith('zone "')]
        if z:
            return 'Zone %s: %s' % (z[0][5:], 'name in the dropdown and on the card' if p[-1] == 'label' else 'description under the zone')
        rest = '.'.join(p[p.index('field "zones"') + 1:])
        return {
            'choiceLabel': 'Small heading above an area with no current zone',
            'placeholder': 'Dropdown prompt (area with no current zone)',
            'empty.title': '"Nothing assigned" screen: title',
            'empty.message': '"Nothing assigned" screen: message',
            'allDone.title': '"Already submitted" screen: title',
            'allDone.message': '"Already submitted" screen: message',
        }.get(rest, rest)
    return j

def where_access(p):
    j = '.'.join(p)
    return {
        'title': 'Page title (in the header band)',
        'back': 'Link above the card back to the form',
        'intro': 'Blue info box at the top (keep the [link text]({var:formUrl}) form)',
        'empty.title': '"No areas assigned" screen: title',
        'empty.message': '"No areas assigned" screen: message',
        'errorContact.subject': 'Subject line of the problem-report email from this page',
    }.get(j, j)

def where_res(path, node_type):
    return path

rows_form = []
walk(form, [], rows_form)
# in the order the page shows them, under group headings
FORM_GROUPS = ['Header', 'Access-list callout (top of the form)', 'Each area (zone cards)', 'Zones: names and descriptions',
               'Bottom of the form', 'Submitting', 'Screens instead of the form', 'Problem-report email']
def form_group(p):
    j = '.'.join(p)
    if p[0] == 'form': return 0
    if 'field "accessNotice"' in p: return 1
    if any(s.startswith('zone "') for s in p): return 3
    if 'field "changeNotice"' in p or 'field "contactNotice"' in p: return 4
    if p[0] == 'submitConfirm' or (p[0] == 'strings' and p[1] in ('submit', 'savingTitle', 'rowsPartial', 'rowsFailed', 'assignAllRequired')): return 5
    if p[0] == 'confirmation' or 'empty' in p or 'allDone' in p: return 6
    if p[0] == 'errorContact': return 7
    return 2  # zones field + the per-area strings
rows_form = sorted(((form_group(p), i, clean(p), where_form(p), en, fr) for i, (p, en, fr) in enumerate(rows_form)))
rows_form = [(FORM_GROUPS[g], key, where, en, fr) for g, i, key, where, en, fr in rows_form]

rows_acc = []
walk(access, [], rows_acc)
rows_acc = [(clean(p), where_access(p), en, fr) for p, en, fr in rows_acc]

# side panel: by group heading and item title, so the SME recognises them
rows_res = []
for gi, g in enumerate(res['sections']):
    gname = g['heading']['en']
    rows_res.append(('group %d › heading' % (gi + 1), 'Side panel: group heading', g['heading']['en'], g['heading'].get('fr', '')))
    for ii, it in enumerate(g['items']):
        kind = {'feature': 'large card', 'link': 'document link', 'text': 'information paragraph', 'contact': 'contact card'}[it['type']]
        for k in ('title', 'text', 'description'):
            if is_pair(it.get(k)):
                what = {'title': 'heading' if it['type'] in ('text', 'contact') else 'title', 'text': 'text', 'description': 'description under it'}[k]
                rows_res.append(('group %d › item %d (%s) › %s' % (gi + 1, ii + 1, it['type'], k),
                                 'Side panel, "%s" › %s %d: %s' % (gname, kind, ii + 1, what), it[k]['en'], it[k].get('fr', '')))

# single values (not translated)
others = [
    ('Form', 'errorContact.email', 'Who users email about a problem (form page)', form['errorContact']['email']),
    ('Access list', 'errorContact.email', 'Who users email about a problem (access list page)', access['errorContact']['email']),
    ('Form', 'form.vars.accessListUrl', 'Address of the access list page (the callout button)', form['form']['vars']['accessListUrl']),
    ('Access list', 'vars.formUrl', 'Address of the form page (the back link)', access['vars']['formUrl']),
]
for gi, g in enumerate(res['sections']):
    for ii, it in enumerate(g['items']):
        if 'url' in it:
            others.append(('Side panel', 'group %d › item %d (%s) › url' % (gi + 1, ii + 1, it['type']),
                           'Link for "%s" (leave blank: shown without a link)' % it['title']['en'], it['url'] or '(blank)'))
        if it['type'] == 'contact':
            others.append(('Side panel', 'group %d › item %d (contact) › name' % (gi + 1, ii + 1), 'Contact card: name', it['name']))
            others.append(('Side panel', 'group %d › item %d (contact) › email' % (gi + 1, ii + 1), 'Contact card: email', it['email']))

# built-in messages: defaults the pages can show, changeable through "strings"
def block(name):
    m = re.search(r'var %s = \{(.*?)\n  \};' % name, js, re.S)
    return dict(re.findall(r"(\w+): '((?:[^'\\]|\\.)*)'", m.group(1)))
D_EN, D_FR = block('DEFAULT_STRINGS'), block('DEFAULT_STRINGS_FR')
BUILTIN_FORM = [
    ('submitting', 'Submit button while saving'),
    ('savingMessage', 'Full-page "saving" overlay: its message'),
    ('assignRetry', '"Try again" button after a load error'),
    ('assignSaved', 'Status tag on an area already saved'),
    ('assignClear', 'Dropdown option that clears the choice'),
    ('assignMore', 'Warning when a person has more areas than one load shows'),
    ('assignKeysMore', 'Warning when too many earlier answers to check'),
    ('assignNoUser', 'Error when the user can\'t be identified'),
    ('pageError', 'Error summary above the submit button'),
    ('submitFailed', 'Error when nothing could be saved'),
    ('errorMailLink', 'Link to email the administrator about an error'),
    ('errorMailAsk', 'First line of that email ("What were you doing…")'),
    ('errorMailInfo', 'Line above the diagnostic details in that email'),
    ('configLoadError', 'Shown if the form can\'t start at all'),
    ('userLoading', 'While the user card loads'),
    ('opensNewTab', 'Read aloud by screen readers on new-tab links'),
]
BUILTIN_ACCESS = [
    ('accessLoading', 'While the access list loads'),
    ('accessLoadError', 'When the access list can\'t be loaded'),
    ('accessNoUser', 'When the user can\'t be identified'),
    ('accessNoEmployees', 'An area with nobody listed'),
    ('accessCount', 'Count beside each area ("5 employee(s)")'),
    ('accessExpandAll', 'Button: expand every area'),
    ('accessCollapseAll', 'Button: collapse every area'),
    ('accessMore', 'Warning when only part of the list could be read'),
    ('accessPartialTitle', 'Screen when none of a partly-read list can be shown: title'),
    ('accessPartialMessage', 'That screen: message'),
]
ov_form, ov_acc = form.get('strings', {}), access.get('strings', {})
builtin_form = [(k, w, D_EN.get(k, ''), D_FR.get(k, '')) for k, w in BUILTIN_FORM if k not in ov_form]
builtin_acc = [(k, w, D_EN.get(k, ''), D_FR.get(k, '')) for k, w in BUILTIN_ACCESS if k not in ov_acc]

# ---------------------------------------------------------------- document
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(27.94), Cm(21.59)  # US Letter landscape
for m in ('left_margin', 'right_margin'): setattr(sec, m, Cm(1.5))
for m in ('top_margin', 'bottom_margin'): setattr(sec, m, Cm(1.5))
st = doc.styles['Normal']; st.font.name = 'Segoe UI'; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Segoe UI')
for h in ('Heading 1', 'Heading 2', 'Heading 3', 'Title'):
    doc.styles[h].font.name = 'Segoe UI'; doc.styles[h].font.color.rgb = RGBColor(0x00, 0x19, 0x28)

GREY, AMBER = '6B7280', 'FFF4CE'
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)
def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)
def no_split(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:cantSplit'); el.set(qn('w:val'), 'true'); trPr.append(el)
def fix_widths(t, widths):
    """Word ignores cell widths under autofit: pin the layout, the grid and every cell"""
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w.twips if hasattr(w, 'twips') else w / 635)))
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = w
def para(cell, text, size=9, bold=False, color=None, italic=False):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text else cell.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(0)
    return p
def is_placeholder(s):
    return '[Placeholder]' in s or '[Texte provisoire]' in s
# the texts the engine runs through prose() (engine 1.4.0+): links, bold,
# italic. Everything else is plain text. Taken from prose()'s call sites in
# src/ps-zone-attestation.js (grep "prose("): intro, headerNote / note text,
# section and page description, field hint and lockNote, zone description,
# empty / allDone / confirmation / submitConfirm / prompt / afterSubmit /
# queryError message, the access intro and empty message, the side panel's
# text items and descriptions (a feature card's own text is plain). Keep in
# step with docs/03-config.md, "Formatting".
PROSE_KEYS = ('intro', 'text', 'description', 'message', 'hint', 'lockNote')
def takes_formatting(key):
    parts = key.split(' › ')
    if parts[-1] not in PROSE_KEYS: return False
    # a note's button label and a feature card's own text are plain; match
    # whole path segments, so a field id like "buttonCount" doesn't count
    return 'button' not in parts and not (len(parts) > 1 and parts[-1] == 'text' and parts[-2].endswith('(feature)'))
if __name__ == '__main__' and os.environ.get('PSZ_SELFTEST'):
    for k, want in [('field "buttonCount" › hint', True), ('field "accessNotice" › button › text', False),
                    ('group 1 › item 1 (feature) › text', False), ('group 1 › item 1 (feature) › description', True),
                    ('field "zones" › zone "Public" › label', False), ('field "x" › lockNote', True)]:
        assert takes_formatting(k) == want, k
    print('takes_formatting: ok'); sys.exit(0)

counter = {'n': 0}
def string_table(prefix, rows, file_label, note_placeholders=True):
    widths = [Cm(1.3), Cm(5.2), Cm(4.9), Cm(4.9), Cm(4.9), Cm(4.9)]
    t = doc.add_table(rows=1, cols=6); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ['#', 'Where it shows', 'Current English', 'Current French', 'New English', 'New French']
    for i, c in enumerate(t.rows[0].cells):
        c.width = widths[i]; para(c, hdr[i], 9, True, 'FFFFFF'); shade(c, '001928')
    repeat_header(t.rows[0])
    group = None
    for i, row in enumerate(rows, 1):
        if len(row) == 5:  # (group, key, where, en, fr): a heading row when the group changes
            g, key, where, en, fr = row
            if g != group:
                group = g
                gr = t.add_row(); no_split(gr)
                merged = gr.cells[0].merge(gr.cells[5])
                para(merged, g, 9.5, True); shade(merged, 'E5EEF6')
        else:
            key, where, en, fr = row
        r = t.add_row(); no_split(r); cells = r.cells
        for j, c in enumerate(cells): c.width = widths[j]
        para(cells[0], '%s-%02d' % (prefix, i), 8, True)
        para(cells[1], where, 9)
        if takes_formatting(key):
            para(cells[1], 'Formatting allowed', 7.5, True, '0F6E3F')
        para(cells[1], key, 7.5, color=GREY, italic=True)
        para(cells[2], en, 9); para(cells[3], fr, 9)
        if note_placeholders and (is_placeholder(en) or is_placeholder(fr)):
            shade(cells[4], AMBER); shade(cells[5], AMBER)
    fix_widths(t, widths)
    doc.add_paragraph()

def value_table(rows):
    widths = [Cm(1.3), Cm(2.6), Cm(7.4), Cm(7.6), Cm(7.2)]
    t = doc.add_table(rows=1, cols=5); t.style = 'Table Grid'
    hdr = ['#', 'File', 'What it is', 'Current value', 'New value']
    for i, c in enumerate(t.rows[0].cells):
        c.width = widths[i]; para(c, hdr[i], 9, True, 'FFFFFF'); shade(c, '001928')
    repeat_header(t.rows[0])
    for i, (f, key, what, val) in enumerate(rows, 1):
        r = t.add_row(); no_split(r); cells = r.cells
        for j, c in enumerate(cells): c.width = widths[j]
        para(cells[0], 'V-%02d' % i, 8, True); para(cells[1], f, 9)
        para(cells[2], what, 9); para(cells[2], key, 7.5, color=GREY, italic=True)
        para(cells[3], val, 9)
        if 'example.com' in val or 'Placeholder' in val or val == '(blank)': shade(cells[4], AMBER)
    fix_widths(t, widths)
    doc.add_paragraph()

doc.add_heading('PS Work Zone Confirmation: wording for review', 0)
p = doc.add_paragraph()
p.add_run('Engine %s · prepared %s · ' % (VERSION, datetime.date.today().isoformat())).font.color.rgb = RGBColor.from_string(GREY)
p.add_run('Every piece of text the form, the employee access list page and the side panel show, in English and French.').font.color.rgb = RGBColor.from_string(GREY)

doc.add_heading('How to fill this in', 1)
for line in [
    'Write new wording in the New English and New French columns. Leave both blank to keep the current text.',
    'Rows shaded yellow still hold placeholder text ([Placeholder] or [Texte provisoire]) and must be filled in before go-live.',
    'Keep anything in curly braces exactly as it is, for example {n}, {total} or {submit}. The page replaces it with a number or a button name.',
    'Keep links in the form [link text](address). You can change the link text in the square brackets; leave the address in round brackets alone unless that row is about it.',
    'Rows marked "Formatting allowed" can also use **bold** and *italic* (two asterisks or one, right up against the words). '
    'Nothing else works: no headings, bullet points, numbered lists or lines. Those show exactly as typed. '
    'In every other row, asterisks and brackets show as typed.',
    'All French text is a draft and needs a French reviewer, even where the English is final.',
    'Zone names and descriptions can change. The zone values the system stores (Public, Public Plus, Standards, Secure Restricted, High Security) can\'t, so only the wording is listed here.',
    'The grey text under "Where it shows" is for the person applying the changes; you can ignore it.',
]:
    doc.add_paragraph(line, style='List Bullet')

doc.add_heading('1. The confirmation form', 1)
doc.add_paragraph('File: src/ps-zone-attestation.json')
string_table('F', rows_form, 'form')

doc.add_heading('2. The employee access list page', 1)
doc.add_paragraph('File: src/ps-zone-attestation.access.json')
string_table('A', rows_acc, 'access')

doc.add_heading('3. The side panel (resources and contact)', 1)
doc.add_paragraph('File: src/ps-zone-attestation.resources.json')
string_table('R', rows_res, 'resources')

doc.add_heading('4. Names, emails and addresses', 1)
doc.add_paragraph('Not translated: one value each. Page addresses only need changing if the SharePoint pages get different names.')
value_table(others)

doc.add_heading('Appendix: built-in messages (optional)', 1)
doc.add_paragraph('Standard messages the pages show in less common situations (errors, loading, buttons). They already have '
                  'English and French wording and don\'t need to change. Fill a row in only if you want different wording.')
doc.add_heading('On the form', 2)
doc.add_paragraph('To change one, it is added to "strings" in src/ps-zone-attestation.json.')
string_table('B', [('strings › ' + k, w, en, fr) for k, w, en, fr in builtin_form], 'builtin', False)
doc.add_heading('On the access list page', 2)
doc.add_paragraph('To change one, it is added to "strings" in src/ps-zone-attestation.access.json.')
string_table('C', [('strings › ' + k, w, en, fr) for k, w, en, fr in builtin_acc], 'builtin', False)

doc.add_heading('For whoever applies the changes (Copilot)', 1)
for line in [
    'Each row\'s grey key is its location in the file named above that table. "›" separates nested names; field "x", zone "x" and section "x" mean the entry with that id or value; "item n" and "group n" count from 1.',
    'Edit the JSON files as text and keep their formatting: a dev script finds exact lines in them. Change only the "en" / "fr" values (and, in section 4, the single values).',
    'Rows in the appendix are not in the files yet: add them to the file\'s "strings" object as "key": { "en": "…", "fr": "…" }.',
    'Then follow docs/05-deploy-steps.md ("Later changes"): edit the deployed copies too. Wording changes need no version bump.',
]:
    doc.add_paragraph(line, style='List Bullet')

doc.save(OUT)
print('rows: form %d, access %d, side panel %d, values %d, built-in %d + %d' % (
    len(rows_form), len(rows_acc), len(rows_res), len(others), len(builtin_form), len(builtin_acc)))
print('placeholders to fill:', sum(1 for r in rows_form + rows_acc + rows_res if is_placeholder(r[2]) or is_placeholder(r[3])))
